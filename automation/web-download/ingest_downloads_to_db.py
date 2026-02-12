#!/usr/bin/env python3
from __future__ import annotations

import hashlib
import json
import sqlite3
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path('/Users/sanghun/.openclaw/workspace/priceit-starter')
DL_DIR = ROOT / 'automation' / 'web-download' / 'downloads'
DB_PATH = ROOT / 'data' / 'priceit.db'
REPORT = ROOT / 'reports' / f"web_download_ingest_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"


def ensure_schema(conn: sqlite3.Connection):
    conn.execute(
        '''
        CREATE TABLE IF NOT EXISTS b2b_download_products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_site TEXT NOT NULL,
            source_file TEXT NOT NULL,
            row_no INTEGER NOT NULL,
            product_name TEXT,
            option_name TEXT,
            tax_type TEXT,
            supply_price REAL,
            retail_price REAL,
            shipping_fee_text TEXT,
            row_hash TEXT NOT NULL,
            ingested_at TEXT NOT NULL,
            UNIQUE(row_hash)
        );
        '''
    )
    conn.execute('CREATE INDEX IF NOT EXISTS idx_b2b_download_products_site ON b2b_download_products(source_site);')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_b2b_download_products_ingested_at ON b2b_download_products(ingested_at);')


def parse_site_from_name(name: str) -> str:
    # ex) 팜허브_1770893240_상품목록.xlsx
    if '_' in name:
        return name.split('_', 1)[0]
    return 'unknown'


def row_hash(site: str, row: tuple) -> str:
    raw = json.dumps([site, *row], ensure_ascii=False, default=str)
    return hashlib.sha1(raw.encode('utf-8')).hexdigest()


def ingest_file(conn: sqlite3.Connection, path: Path):
    site = parse_site_from_name(path.name)
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    inserted = 0
    skipped = 0
    ts = datetime.now().isoformat(timespec='seconds')

    for i in range(2, ws.max_row + 1):
        product_name = ws.cell(i, 3).value
        option_name = ws.cell(i, 4).value
        tax_type = ws.cell(i, 5).value
        supply_price = ws.cell(i, 6).value
        retail_price = ws.cell(i, 7).value
        shipping_fee_text = ws.cell(i, 10).value

        if not product_name:
            continue

        h = row_hash(site, (product_name, option_name, tax_type, supply_price, retail_price, shipping_fee_text))
        cur = conn.execute(
            '''
            INSERT OR IGNORE INTO b2b_download_products
            (source_site, source_file, row_no, product_name, option_name, tax_type, supply_price, retail_price, shipping_fee_text, row_hash, ingested_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''',
            (site, str(path), i, product_name, option_name, tax_type, supply_price, retail_price, shipping_fee_text, h, ts),
        )
        if cur.rowcount:
            inserted += 1
        else:
            skipped += 1

    return {
        'file': str(path),
        'site': site,
        'inserted': inserted,
        'skipped': skipped,
        'rows': ws.max_row - 1,
    }


def main():
    files = sorted(DL_DIR.glob('*.xlsx'))
    conn = sqlite3.connect(DB_PATH)
    try:
        ensure_schema(conn)
        results = []
        for f in files:
            results.append(ingest_file(conn, f))
        conn.commit()

        summary = {
            'download_dir': str(DL_DIR),
            'db_path': str(DB_PATH),
            'processed_files': len(files),
            'results': results,
        }
        REPORT.parent.mkdir(parents=True, exist_ok=True)
        REPORT.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        print(f'Report: {REPORT}')
    finally:
        conn.close()


if __name__ == '__main__':
    main()
